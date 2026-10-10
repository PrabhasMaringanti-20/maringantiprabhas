"""
Phase 13 - 10-slide executive presentation (presentation/NorthStar_Executive_Presentation.pptx).
Uses the charts from scripts/04_build_charts.py. Every number is taken from sql/README.md / insights.

Run:  python scripts/06_build_presentation.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "screenshots"
OUT = ROOT / "presentation" / "NorthStar_Executive_Presentation.pptx"
OUT.parent.mkdir(exist_ok=True)

INK, INK2, BLUE, ORANGE, PANEL = (RGBColor(0x0B, 0x0B, 0x0B), RGBColor(0x52, 0x51, 0x4E), RGBColor(0x2A, 0x78, 0xD6),
                                  RGBColor(0xEB, 0x68, 0x34), RGBColor(0xF1, 0xF0, 0xEC))
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
W = prs.slide_width


def text(slide, x, y, w, h, content, size=16, bold=False, color=INK, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    items = content if isinstance(content, list) else [content]
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        run = p.add_run()
        run.text = item
        run.font.size, run.font.bold, run.font.color.rgb, run.font.name = Pt(size), bold, color, "Calibri"
        p.space_after = Pt(8)
    return tb


def base(title, kicker, n):
    s = prs.slides.add_slide(BLANK)
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, W, Inches(0.12))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    text(s, 0.6, 0.35, 12, 0.4, kicker.upper(), 12, True, BLUE)
    text(s, 0.6, 0.7, 12.2, 0.9, title, 28, True)
    text(s, 0.6, 7.0, 9, 0.3, "NorthStar Retail UK (fictional) · simulated data · portfolio project", 10, False, INK2)
    text(s, 12.2, 7.0, 0.6, 0.3, str(n), 10, False, INK2, PP_ALIGN.RIGHT)
    return s


def panel(s, x, y, w, h):
    r = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid()
    r.fill.fore_color.rgb = PANEL
    r.line.fill.background()
    return r


def kpi(s, x, y, label, value, sub):
    panel(s, x, y, 2.35, 1.25)
    text(s, x + 0.15, y + 0.08, 2.1, 0.3, label, 12, False, INK2)
    text(s, x + 0.15, y + 0.38, 2.1, 0.5, value, 26, True)
    text(s, x + 0.15, y + 0.85, 2.1, 0.3, sub, 11, False, INK2)


def pic(s, name, x, y, w=None, h=None):
    kw = {}
    if w:
        kw["width"] = Inches(w)
    if h:
        kw["height"] = Inches(h)
    s.shapes.add_picture(str(IMG / name), Inches(x), Inches(y), **kw)


def notes(s, txt):
    s.notes_slide.notes_text_frame.text = txt


# 1 Problem
s = base("Revenue is growing. Profit isn't keeping up.", "1 · The problem", 1)
text(s, 0.6, 1.8, 6.2, 4.5, [
    "Leadership receives four disconnected weekly reports (stores, e-commerce, finance, returns) that disagree.",
    "Revenue grew in 2025, but no one could say where profit was being made or lost.",
    "Suspected causes: weak-margin categories, heavy discounting, rising returns, under-performing stores.",
    "This project builds one trusted source of truth and answers: what is happening, why, where, and what to do.",
], 17)
kpi(s, 7.4, 1.9, "Revenue 2025", "£4.58M", "+15.7% vs 2024")
kpi(s, 10.0, 1.9, "Gross profit 2025", "£1.88M", "+9.4% vs 2024")
kpi(s, 7.4, 3.4, "Gross margin", "41.0%", "−2.4 pts vs 2024")
kpi(s, 10.0, 3.4, "Return rate", "8.3%", "+1.2 pts vs 2024")
notes(s, "Open with the tension: growth looks healthy, but profit grew at less than two-thirds the rate of revenue.")

# 2 Objectives
s = base("What leadership asked for", "2 · Business objectives", 2)
rows = [("O1", "One trusted source of truth", "KPIs reconcile SQL ↔ Power BI"),
        ("O2", "See where profit is made or lost", "Drill company → region → store → product"),
        ("O3", "Quantify the cost of discounting", "Margin by discount band, break-even discount"),
        ("O4", "Explain rising returns", "By category, channel, brand and reason"),
        ("O5", "Understand customer value", "Segments, retention, RFM"),
        ("O6", "Recommend prioritised actions", "≥ 8 recommendations with stated assumptions")]
tbl = s.shapes.add_table(len(rows) + 1, 3, Inches(0.6), Inches(1.8), Inches(12), Inches(4.2)).table
for j, h in enumerate(["#", "Objective", "Measure of success"]):
    tbl.cell(0, j).text = h
for i, r in enumerate(rows, 1):
    for j, v in enumerate(r):
        tbl.cell(i, j).text = v
tbl.columns[0].width, tbl.columns[1].width, tbl.columns[2].width = Inches(0.8), Inches(5.2), Inches(6.0)
for row in tbl.rows:
    for c in row.cells:
        for p in c.text_frame.paragraphs:
            for r_ in p.runs:
                r_.font.size = Pt(15)
text(s, 0.6, 6.25, 12, 0.5, "Out of scope: inventory, marketing spend, overheads → we report gross profit, not net profit.", 14, False, INK2)

# 3 Data
s = base("185,858 order lines, 24 months, one star schema", "3 · Data & approach", 3)
text(s, 0.6, 1.8, 5.6, 4.8, [
    "Pipeline: raw CSV extracts → Python validation & cleaning → PostgreSQL star schema → 40 SQL business questions → Power BI model + DAX → insights.",
    "Scale: 89,528 orders · 9,530 known customers · 1,220 products · 24 stores + online · 12 UK regions · Jan 2024 – Dec 2025.",
    "Model: FactSales (grain = order line) + FactReturns, sharing six dimensions (Date, Customer, Product, Store, Region, Channel).",
    "Synthetic data with designed business patterns: seasonality, channel shift, discounting, supplier issues.",
], 15)
panel(s, 6.6, 1.8, 6.1, 4.8)
text(s, 6.8, 1.9, 5.8, 4.6, [
    "            DimDate      DimProduct",
    "                  \\           /",
    "DimCustomer — FactSales — DimStore",
    "                  /           \\",
    "         DimChannel     DimRegion",
    "",
    "FactReturns shares the same dimensions",
    "(return date active; order date inactive",
    "→ USERELATIONSHIP in DAX)",
], 15, False, INK)

# 4 Data quality
s = base("Messy sources, measured and fixed, not hidden", "4 · Data quality", 4)
dq = [("Check", "Before", "After"),
      ("Sales validity (date, qty, discount, channel)", "91.7%", "100%"),
      ("Channel labels (target 4)", "14", "4"),
      ("Category labels (target 8)", "28", "8"),
      ("Duplicate sales rows", "900", "0"),
      ("Duplicate customers (same email)", "180", "0"),
      ("Returns linked to a sale", "99.87%", "100%")]
tbl = s.shapes.add_table(len(dq), 3, Inches(0.6), Inches(1.8), Inches(7.4), Inches(3.8)).table
for i, r in enumerate(dq):
    for j, v in enumerate(r):
        tbl.cell(i, j).text = v
        for p in tbl.cell(i, j).text_frame.paragraphs:
            for r_ in p.runs:
                r_.font.size = Pt(14)
tbl.columns[0].width, tbl.columns[1].width, tbl.columns[2].width = Inches(4.8), Inches(1.3), Inches(1.3)
text(s, 8.5, 1.8, 4.3, 4.8, [
    "40 documented rules: fix if certain, flag if uncertain, quarantine if untrustworthy.",
    "90 sales rows quarantined (0.05%): impossible dates, negative quantities, unknown SKUs.",
    "Guest (blank) customers kept as a legitimate 'unknown member', not deleted.",
    "Proof: 100% of cleaned lines match the answer key; 13/13 SQL checks pass.",
], 15)

# 5 Sales
s = base("Growth is digital; stores are flat like-for-like", "5 · Sales", 5)
pic(s, "04_channel_growth.png", 0.5, 1.7, w=6.3)
pic(s, "01_monthly_revenue_margin.png", 7.0, 1.7, w=5.9)
text(s, 0.6, 5.0, 6.2, 1.8, [
    "Mobile App +49%, Website +22%, Click & Collect +24%, In-Store +4%.",
    "Like-for-like stores +5.3% vs online hub +29%. Headline growth flatters the store estate.",
], 14)

# 6 Profitability
s = base("Margin fell 2.4 pts: discounting and unrecovered cost inflation", "6 · Profitability", 6)
pic(s, "09_margin_bridge.png", 0.5, 1.7, h=3.85)
pic(s, "03_discount_vs_margin.png", 6.9, 1.7, h=3.85)
text(s, 0.6, 5.7, 12.2, 1.3, [
    "Electronics: 19% of revenue, 6% of profit. 2025 revenue +14.6% but gross profit −13%. Breaks even at an 18.7% discount.",
    "Lines discounted 30%+ in 2025: £406K revenue, only £56K gross profit.",
], 14)

# 7 Customers
s = base("A loyal core, and 1,450 valuable customers drifting away", "7 · Customers", 7)
pic(s, "10_rfm_segments.png", 0.5, 1.7, w=6.6)
text(s, 7.4, 1.9, 5.4, 4.8, [
    "14% of customers generate 50% of known-customer revenue; 38% generate 80%.",
    "Returning customers place ~93% of orders.",
    "Premium: £1,380 per customer, buys every ~20 days (median).",
    "'At risk' segment: formerly loyal (avg 7.7 orders), no purchase for ~8 months, 14.6% of revenue.",
    "32–34% of in-store revenue is anonymous: can't be retained or analysed.",
], 15)

# 8 Returns
s = base("Online returns rising, driven by one footwear supplier", "8 · Returns", 8)
pic(s, "06_online_apparel_return_trend.png", 0.5, 1.7, w=6.4)
pic(s, "07_brand_return_rates.png", 7.0, 1.7, w=5.9)
text(s, 0.6, 4.75, 12.2, 2.0, [
    "Online clothing & footwear return rate ~16% → ~27%; refunds £83K → £160K.",
    "Stride Lab 10% → 28% (83% 'wrong size / fit'). Voltix smart home: 22% vs 6% for other brands ('faulty').",
    "After refunds, website margin is 28% vs 39% in store.",
], 14)

# 9 Key insights
s = base("Five things leadership should know", "9 · Key insights", 9)
ins = [("1", "Profit grows slower than sales", "Margin −2.4 pts: −1.6 cost inflation not passed on, −0.8 more discounting."),
       ("2", "Electronics sells volume, not profit", "11% margin in 2025; discounts over ~19% sell below cost."),
       ("3", "Deep discounts don't pay", "30%+ discounts: 14% margin on £406K of 2025 revenue."),
       ("4", "Returns are a supplier & sizing problem", "Stride Lab sizing and Voltix faults explain most of the spike."),
       ("5", "Two stores discount like outlets", "Newcastle 15% average discount vs ~5% at peers → 37% margin vs ~44%.")]
for i, (n_, h, b) in enumerate(ins):
    y = 1.75 + i * 1.02
    panel(s, 0.6, y, 12.1, 0.88)
    text(s, 0.75, y + 0.12, 0.6, 0.6, n_, 26, True, BLUE)
    text(s, 1.45, y + 0.08, 4.6, 0.7, h, 17, True)
    text(s, 6.0, y + 0.14, 6.6, 0.7, b, 15, False, INK2)

# 10 Recommendations
s = base("Start with the quick wins: ~£120–140K gross profit a year", "10 · Recommendations & next steps", 10)
pic(s, "11_impact_effort_matrix.png", 0.4, 1.6, h=5.2)
text(s, 7.3, 1.7, 5.6, 5.2, [
    "30 days: discount guardrails by category (Electronics max 15%), ~£28K;",
    "Stride Lab & Voltix supplier action, ~£30–49K; Newcastle promotion sign-off, ~£6K.",
    "60 days: size guidance on top apparel lines (~£20K per 3 pts of returns); win-back campaign, ~£13K.",
    "90 days: Electronics price/cost review, ~£25K; till customer capture; report margin after returns per channel.",
    "Estimates are illustrative, based on simulated data with stated assumptions (insights/02_recommendations.md).",
], 14)

prs.save(OUT)
print("wrote", OUT)
